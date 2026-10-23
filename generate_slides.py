"""
Generate a 5-slide PowerPoint presentation for the Revenue Intelligence Inbox submission.
Run: python generate_slides.py
Requires: pip install python-pptx
"""

from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN
import os


def rgb(r, g, b):
    return RGBColor(r, g, b)


BG_DARK = rgb(10, 14, 26)
ACCENT_BLUE = rgb(96, 165, 250)
ACCENT_PURPLE = rgb(167, 139, 250)
ACCENT_GREEN = rgb(52, 211, 153)
ACCENT_AMBER = rgb(251, 191, 36)
ACCENT_RED = rgb(248, 113, 113)
TEXT_MAIN = rgb(226, 232, 240)
TEXT_MUTED = rgb(100, 116, 139)
TEXT_BRIGHT = rgb(255, 255, 255)


def set_slide_bg(slide, prs):
    from pptx.util import Emu
    fill = slide.background.fill
    fill.solid()
    fill.fore_color.rgb = BG_DARK


def add_textbox(slide, text, left, top, width, height,
                font_size=14, bold=False, color=TEXT_MAIN,
                align=PP_ALIGN.LEFT, italic=False, wrap=True):
    txBox = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    tf = txBox.text_frame
    tf.word_wrap = wrap
    p = tf.paragraphs[0]
    p.alignment = align
    run = p.add_run()
    run.text = text
    run.font.size = Pt(font_size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = color
    run.font.name = "Calibri"
    return txBox


def add_bullet_box(slide, items, left, top, width, height,
                   font_size=13, color=TEXT_MAIN, bullet_color=ACCENT_BLUE):
    txBox = slide.shapes.add_textbox(
        Inches(left), Inches(top), Inches(width), Inches(height)
    )
    tf = txBox.text_frame
    tf.word_wrap = True
    for i, item in enumerate(items):
        if i == 0:
            p = tf.paragraphs[0]
        else:
            p = tf.add_paragraph()
        p.space_before = Pt(4)
        # Bullet dot
        run0 = p.add_run()
        run0.text = "▸  "
        run0.font.size = Pt(font_size)
        run0.font.color.rgb = bullet_color
        run0.font.name = "Calibri"
        # Item text
        run1 = p.add_run()
        run1.text = item
        run1.font.size = Pt(font_size)
        run1.font.color.rgb = color
        run1.font.name = "Calibri"


def add_divider(slide, top, width=9.2, left=0.4, color=ACCENT_BLUE):
    from pptx.util import Inches, Pt
    from pptx.oxml.ns import qn
    from lxml import etree
    line = slide.shapes.add_shape(
        1,  # MSO_SHAPE_TYPE.LINE is not directly accessible this way
        Inches(left), Inches(top), Inches(width), Inches(0.01)
    )
    line.fill.solid()
    line.fill.fore_color.rgb = color
    line.line.color.rgb = color


def slide_1(prs):
    """Slide 1: Manager Problem & Main Findings"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])  # blank
    set_slide_bg(slide, prs)

    add_textbox(slide, "01", 0.4, 0.2, 1.5, 0.5,
                font_size=28, bold=True, color=ACCENT_PURPLE)
    add_textbox(slide, "THE MANAGER PROBLEM & MAIN FINDINGS",
                1.6, 0.25, 8, 0.5, font_size=12, color=TEXT_MUTED)
    add_textbox(slide, "Revenue Intelligence Inbox — SpendNest",
                0.4, 0.7, 9, 0.6, font_size=20, bold=True, color=TEXT_BRIGHT)

    add_textbox(slide, "The Problem", 0.4, 1.45, 4, 0.4,
                font_size=14, bold=True, color=ACCENT_BLUE)
    problem_items = [
        "A sales manager gets calls, meetings and transcripts from several systems.",
        "Records arrive late (M07 transcript delayed 1 day).",
        "Stored labels may not match what was actually said.",
        "No single view connects activity → analysis → action.",
    ]
    add_bullet_box(slide, problem_items, 0.4, 1.85, 4.5, 2.0, font_size=11)

    add_textbox(slide, "What I Built", 5.2, 1.45, 4.2, 0.4,
                font_size=14, bold=True, color=ACCENT_GREEN)
    built_items = [
        "44-record inbox with search, filter and routing.",
        "Import system handling duplicate-safe, revision-ordered updates.",
        "Pre-generated evidence-based analysis for 18 transcripts.",
        "Manager brief for 21 Aug 2026 with priority actions.",
        "Export archive with machine-readable manifest.",
    ]
    add_bullet_box(slide, built_items, 5.2, 1.85, 4.5, 2.0,
                   font_size=11, bullet_color=ACCENT_GREEN)

    add_textbox(slide, "Three Key Findings", 0.4, 4.0, 9, 0.4,
                font_size=14, bold=True, color=ACCENT_AMBER)

    findings = [
        ("M03 Turn 23", "Pricing discrepancy mid-demo (INR 26k vs 39k) → Sales Coaching route flagged."),
        ("M04 Turn 11", "Cedarwave's core need (credit cards) cannot be met by SpendNest prepaid → Human Review."),
        ("M09 Turn 21", "Brookmint reset token expired immediately on fresh send → Engineering bug ticket."),
    ]
    for i, (ref, desc) in enumerate(findings):
        y = 4.45 + i * 0.55
        add_textbox(slide, ref, 0.4, y, 1.5, 0.45,
                    font_size=10, bold=True, color=ACCENT_AMBER)
        add_textbox(slide, desc, 1.9, y, 7.5, 0.45, font_size=10, color=TEXT_MAIN)

    add_textbox(slide, "SpendNest Revenue Intelligence Inbox  ·  21 August 2026",
                0.4, 7.1, 9, 0.3, font_size=9, color=TEXT_MUTED, align=PP_ALIGN.CENTER)


def slide_2(prs):
    """Slide 2: Interface Choices & What Was Left Out"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, prs)

    add_textbox(slide, "02", 0.4, 0.2, 1.5, 0.5, font_size=28, bold=True, color=ACCENT_BLUE)
    add_textbox(slide, "INTERFACE CHOICES & DELIBERATE OMISSIONS",
                1.6, 0.25, 8, 0.5, font_size=12, color=TEXT_MUTED)
    add_textbox(slide, "Design Decisions",
                0.4, 0.7, 9, 0.5, font_size=20, bold=True, color=TEXT_BRIGHT)

    add_textbox(slide, "Included", 0.4, 1.4, 4, 0.4,
                font_size=14, bold=True, color=ACCENT_GREEN)
    included = [
        "5-page app: Dashboard, Inbox, Import, Export, Brief.",
        "KPI cards: 44 records, 30 calls, 14 meetings, transcript coverage.",
        "Per-record detail: metadata + transcript/turns + analysis + route editor.",
        "Duplicate-safe import with revision ordering (no older rev overwrites newer).",
        "Manual route override with timestamp, note, and visual distinction from suggested.",
        "Stale analysis warning when transcript changes without re-analysis.",
        "Machine-readable manifest (JSON + CSV) and downloadable ZIP.",
        "Denominator notes and source-reported labels on every metric.",
    ]
    add_bullet_box(slide, included, 0.4, 1.8, 4.5, 3.5, font_size=10, bullet_color=ACCENT_GREEN)

    add_textbox(slide, "Deliberately Left Out", 5.2, 1.4, 4.2, 0.4,
                font_size=14, bold=True, color=ACCENT_RED)
    omitted = [
        "Pipeline value totals — no authoritative CRM data supplied.",
        "Cross-person performance rankings — selected sample only.",
        "Live AI API calls — pre-generated analysis is transparent and sufficient.",
        "Date filter — all activity is on one day; date filter adds no value.",
        "Account matching across records — no verified company-ID join.",
        "Invented summaries for no-transcript records — shows 'no text' honestly.",
        "Auto-generated AI summaries for new transcripts — would need API key.",
    ]
    add_bullet_box(slide, omitted, 5.2, 1.8, 4.5, 3.0, font_size=10,
                   bullet_color=ACCENT_RED, color=TEXT_MAIN)

    add_textbox(slide, "Routing Badge Legend",
                0.4, 5.4, 9, 0.4, font_size=12, bold=True, color=TEXT_MUTED)
    badges = [
        ("Sales Coaching", ACCENT_AMBER),
        ("Deal Follow-up", ACCENT_GREEN),
        ("Customer Follow-up", ACCENT_BLUE),
        ("Internal/Vendor", ACCENT_PURPLE),
        ("Human Review", ACCENT_RED),
    ]
    for i, (name, color) in enumerate(badges):
        add_textbox(slide, f"■ {name}", 0.4 + i * 1.9, 5.8, 1.8, 0.35,
                    font_size=10, color=color, bold=True)

    add_textbox(slide, "SpendNest Revenue Intelligence Inbox  ·  21 August 2026",
                0.4, 7.1, 9, 0.3, font_size=9, color=TEXT_MUTED, align=PP_ALIGN.CENTER)


def slide_3(prs):
    """Slide 3: Routing Design & One Difficult Case"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, prs)

    add_textbox(slide, "03", 0.4, 0.2, 1.5, 0.5, font_size=28, bold=True, color=ACCENT_GREEN)
    add_textbox(slide, "ROUTING DESIGN & ONE DIFFICULT CASE",
                1.6, 0.25, 8, 0.5, font_size=12, color=TEXT_MUTED)
    add_textbox(slide, "Routing & Folder Design",
                0.4, 0.7, 9, 0.5, font_size=20, bold=True, color=TEXT_BRIGHT)

    add_textbox(slide, "Routing Taxonomy", 0.4, 1.35, 4, 0.4,
                font_size=14, bold=True, color=ACCENT_BLUE)
    taxonomy = [
        "Sales Coaching — Evidence of rep behaviour to address (e.g., C01 disclosure failure, M03 pricing).",
        "Deal Follow-up — Active opportunity with a clear next step (e.g., M01 MoU, C03 email).",
        "Customer Follow-up — Existing customer with open issues or pending decisions.",
        "Internal/Vendor — Internal process or vendor integration discussion (e.g., M02, M10).",
        "Human Review — No usable transcript; or multiple unresolved issues needing a human decision.",
    ]
    add_bullet_box(slide, taxonomy, 0.4, 1.75, 4.5, 3.0, font_size=10)

    add_textbox(slide, "Folder Structure", 5.2, 1.35, 4.2, 0.4,
                font_size=14, bold=True, color=ACCENT_PURPLE)
    folder_text = (
        "exports/\n"
        "├── Deal_Follow-up/\n"
        "├── Sales_Coaching/\n"
        "├── Customer_Follow-up/\n"
        "├── Internal-Vendor/\n"
        "├── Human_Review/\n"
        "├── manifest.json\n"
        "├── manifest.csv\n"
        "└── revenue_intelligence_export.zip"
    )
    add_textbox(slide, folder_text, 5.2, 1.75, 4.2, 2.2,
                font_size=10, color=ACCENT_GREEN)

    add_textbox(slide, "Difficult Case: M07 — Willowfield Trust",
                0.4, 4.8, 9, 0.45, font_size=14, bold=True, color=ACCENT_AMBER)

    difficult = [
        "M07 was in the initial batch with an empty turns array and unavailable_reason: 'Transcript not included; later batch supplies it.'",
        "Challenge: M07 needed a route before the transcript arrived — but routing without evidence would be invented analysis.",
        "Decision: Route M07 to Human Review initially with the unavailable_reason shown clearly. No summary invented.",
        "After update batch (revision 2): M07 receives its transcript (27 turns). Analysis is marked stale_pending_review.",
        "The manager sees a warning banner and must update the route manually — making the update visible rather than silent.",
        "Manual override (set before the update) survives intact. The new transcript is inspectable in the turns panel.",
    ]
    add_bullet_box(slide, difficult, 0.4, 5.25, 9.2, 2.1, font_size=10)

    add_textbox(slide, "SpendNest Revenue Intelligence Inbox  ·  21 August 2026",
                0.4, 7.1, 9, 0.3, font_size=9, color=TEXT_MUTED, align=PP_ALIGN.CENTER)


def slide_4(prs):
    """Slide 4: Validation — numbers and analytical conclusion"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, prs)

    add_textbox(slide, "04", 0.4, 0.2, 1.5, 0.5, font_size=28, bold=True, color=ACCENT_AMBER)
    add_textbox(slide, "VALIDATION — NUMBERS & ANALYTICAL CONCLUSION",
                1.6, 0.25, 8, 0.5, font_size=12, color=TEXT_MUTED)
    add_textbox(slide, "Checking a Key Number",
                0.4, 0.7, 9, 0.5, font_size=20, bold=True, color=TEXT_BRIGHT)

    add_textbox(slide, "Number Verified: 18 Usable Transcripts After Update",
                0.4, 1.35, 9, 0.4, font_size=14, bold=True, color=ACCENT_GREEN)
    number_steps = [
        "Step 1: Loaded 01_initial.json — python test_imports.py confirms 44 records, 17 usable texts.",
        "Step 2: 10 calls with non-null transcript: C01–C10 (verified by iterating payload.transcript).",
        "Step 3: 7 meetings with ≥1 non-blank turn: M01–M04, M06, M09, M10 (M07 had turns=[] in initial).",
        "Step 4: Loaded 02_update.json (revision 2) — M07 received 27 turns. is_usable_transcript() now True.",
        "Step 5: Count recalculated → 18 usable. Total records still 44 (no new record created for M07).",
        "Verified: test_imports.py asserts all conditions; all assertions pass.",
    ]
    add_bullet_box(slide, number_steps, 0.4, 1.8, 9.2, 2.2, font_size=10, bullet_color=ACCENT_GREEN)

    add_textbox(slide, "Analytical Conclusion Verified: M03 Pricing Discrepancy",
                0.4, 4.1, 9, 0.4, font_size=14, bold=True, color=ACCENT_AMBER)
    analysis_steps = [
        "Claim: Rep quoted INR 26,000 at an earlier meeting but the demo showed INR 39,000.",
        "Source: M03 Turn 23 — Customer A: 'But you previously told me INR 26,000. Why has that become INR 39,000?'",
        "Source: M03 Turn 24 — Rep: 'INR 26,000 plus tax is the recurring subscription, and INR 13,000 is the one-time setup.'",
        "Conclusion is grounded in transcript text, not assumed. The discrepancy is a customer-raised challenge, not a rep error alone.",
        "What the evidence does NOT establish: whether the prior quote was verbal or written; whether the customer will proceed.",
        "Product effect: M03 routes to both Deal Follow-up AND Sales Coaching. Coaching item in Manager Brief.",
    ]
    add_bullet_box(slide, analysis_steps, 0.4, 4.55, 9.2, 2.1, font_size=10, bullet_color=ACCENT_AMBER)

    add_textbox(slide, "SpendNest Revenue Intelligence Inbox  ·  21 August 2026",
                0.4, 7.1, 9, 0.3, font_size=9, color=TEXT_MUTED, align=PP_ALIGN.CENTER)


def slide_5(prs):
    """Slide 5: Trade-offs, AI usage, limitations, improvements"""
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    set_slide_bg(slide, prs)

    add_textbox(slide, "05", 0.4, 0.2, 1.5, 0.5, font_size=28, bold=True, color=ACCENT_RED)
    add_textbox(slide, "TRADE-OFFS · AI USAGE · LIMITATIONS · NEXT",
                1.6, 0.25, 8, 0.5, font_size=12, color=TEXT_MUTED)
    add_textbox(slide, "Decisions, Limitations & What's Next",
                0.4, 0.7, 9, 0.5, font_size=20, bold=True, color=TEXT_BRIGHT)

    add_textbox(slide, "Rejected Decision", 0.4, 1.35, 4.5, 0.4,
                font_size=14, bold=True, color=ACCENT_RED)
    rejected = [
        "Considered: Auto-generate summaries via live LLM API on new transcript arrival.",
        "Rejected: Requires API key (assignment says not required); LLMs can hallucinate transcript details.",
        "Accepted trade-off: Pre-generated analysis is static. New transcripts mark record stale.",
        "Product choice: stale_pending_review flag + warning banner + manual override editor.",
    ]
    add_bullet_box(slide, rejected, 0.4, 1.75, 4.5, 1.8, font_size=10, bullet_color=ACCENT_RED)

    add_textbox(slide, "AI Usage", 5.2, 1.35, 4.2, 0.4,
                font_size=14, bold=True, color=ACCENT_BLUE)
    ai_items = [
        "Claude used to write application code.",
        "All analysis, route decisions, evidence quotes written by candidate after reading source transcripts.",
        "No live AI API called at runtime.",
        "Pre-generated content clearly identified in README and in-app.",
    ]
    add_bullet_box(slide, ai_items, 5.2, 1.75, 4.2, 1.8, font_size=10, bullet_color=ACCENT_BLUE)

    add_textbox(slide, "Known Limitations", 0.4, 3.65, 4.5, 0.4,
                font_size=14, bold=True, color=ACCENT_AMBER)
    limitations = [
        "B01–B20 (20 background calls): no transcript, no analysis beyond 'background call'.",
        "M05/M08/M11–M14: recording cancelled or not scheduled; no turns.",
        "Single state.json file — no multi-user concurrency.",
        "No authentication; local review use only.",
    ]
    add_bullet_box(slide, limitations, 0.4, 4.05, 4.5, 1.8, font_size=10, bullet_color=ACCENT_AMBER)

    add_textbox(slide, "What I Would Improve Next", 5.2, 3.65, 4.2, 0.4,
                font_size=14, bold=True, color=ACCENT_GREEN)
    improvements = [
        "Optional LLM integration: human-confirmed re-analysis after transcript update.",
        "SQLite backend for concurrent users and query performance.",
        "Webhook endpoint to receive live batch deliveries.",
        "Richer coaching views: per-rep trend across multiple days.",
        "Screen recording workflow guide for managers.",
    ]
    add_bullet_box(slide, improvements, 5.2, 4.05, 4.2, 1.8, font_size=10, bullet_color=ACCENT_GREEN)

    add_textbox(slide, "Time Spent: ~8 hours  ·  Data reading: 1h  ·  Coding: 5h  ·  Analysis writing: 1.5h  ·  Docs: 0.5h",
                0.4, 6.5, 9.2, 0.35, font_size=10, color=TEXT_MUTED, align=PP_ALIGN.CENTER)
    add_textbox(slide, "SpendNest Revenue Intelligence Inbox  ·  21 August 2026",
                0.4, 7.1, 9, 0.3, font_size=9, color=TEXT_MUTED, align=PP_ALIGN.CENTER)


def main():
    try:
        from pptx import Presentation
    except ImportError:
        print("python-pptx not installed. Run: pip install python-pptx")
        return

    prs = Presentation()
    prs.slide_width = Inches(10)
    prs.slide_height = Inches(7.5)

    slide_1(prs)
    slide_2(prs)
    slide_3(prs)
    slide_4(prs)
    slide_5(prs)

    out_path = "Revenue_Intelligence_Inbox_Slides.pptx"
    prs.save(out_path)
    print(f"Saved: {out_path}")


if __name__ == "__main__":
    main()
